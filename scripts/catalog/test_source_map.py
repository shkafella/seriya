from __future__ import annotations

import copy
import importlib.util
import sys
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook


SCRIPT_PATH = Path(__file__).with_name("build-seed.py")
SPEC = importlib.util.spec_from_file_location("catalog_build_seed", SCRIPT_PATH)
if SPEC is None or SPEC.loader is None:
    raise RuntimeError(f"Cannot load {SCRIPT_PATH}")
BUILD_SEED = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = BUILD_SEED
SPEC.loader.exec_module(BUILD_SEED)


def row(sheet: str, row_number: int, name: str, sku: str) -> dict[str, object]:
    return {
        "sourceSheet": sheet,
        "sourceRow": row_number,
        "sourceName": name,
        "sourceSku": sku,
    }


class SourceMapIdentityTests(unittest.TestCase):
    def setUp(self) -> None:
        self.rows = [
            row("низ гола", 10, "Нижний модуль 400", "НГП 400"),
            row("низ гола", 11, "Нижний модуль 450", "НГП 450"),
        ]
        self.source_map = {
            "schemaVersion": 1,
            "spreadsheetId": BUILD_SEED.SPREADSHEET_ID,
            "records": [],
        }
        BUILD_SEED.assign_sheet_source_record_ids(self.rows, self.source_map, True)

    def remap(self, changed_rows: list[dict[str, object]]) -> list[str]:
        source_map = copy.deepcopy(self.source_map)
        BUILD_SEED.assign_sheet_source_record_ids(changed_rows, source_map, False)
        return [changed_row["sourceRecordId"] for changed_row in changed_rows]

    def test_sku_change_preserves_source_record_id(self) -> None:
        changed = copy.deepcopy(self.rows)
        changed[0]["sourceSku"] = "НГП 401"

        self.assertEqual(
            self.remap(changed)[0],
            self.rows[0]["sourceRecordId"],
        )

    def test_row_move_preserves_source_record_id(self) -> None:
        changed = copy.deepcopy(self.rows)
        changed[1]["sourceRow"] = 30

        self.assertEqual(
            self.remap(changed)[1],
            self.rows[1]["sourceRecordId"],
        )

    def test_simultaneous_move_and_rename_requires_explicit_mapping(self) -> None:
        changed = copy.deepcopy(self.rows)
        changed[0]["sourceRow"] = 30
        changed[0]["sourceSku"] = "НГП 401"

        with self.assertRaisesRegex(ValueError, "sourceRecordId"):
            self.remap(changed)

    def test_derived_record_id_is_reused(self) -> None:
        first = BUILD_SEED.assign_derived_source_record_id(
            self.source_map,
            "wall-drying-aventos-450",
            True,
        )
        second = BUILD_SEED.assign_derived_source_record_id(
            self.source_map,
            "wall-drying-aventos-450",
            False,
        )

        self.assertEqual(first, second)


class HardwareSpecificationTests(unittest.TestCase):
    def test_one_and_two_front_modules_keep_base_and_extra_hinges(self) -> None:
        hinge = {
            "row": 8,
            "item": "Петля накладная + ответка крест",
            "quantity": 2,
            "unit": "шт",
            "external": True,
        }
        for family_id in BUILD_SEED.TWO_FRONT_FAMILIES:
            family = BUILD_SEED.FAMILIES[family_id]
            rows_by_sheet = {
                (family.hardware_source, family.hardware_sheet): [hinge],
            }
            for front_count, expected_hinges in ((1, 2), (2, 4)):
                module = {
                    "dimensionsMm": {"width": 600},
                    "variant": {"frontCount": front_count},
                }
                _, supplied, _ = BUILD_SEED.calculate_hardware(
                    family,
                    module,
                    rows_by_sheet,
                    {"Саморез 4х16мм": Decimal("1")},
                )

                self.assertEqual(
                    supplied[0]["quantity"],
                    expected_hinges,
                    f"{family_id=} {front_count=}",
                )

    def test_fridge_keeps_common_customer_items_and_side_panel_extras(self) -> None:
        base = {"row": 2, "item": "Евровинт", "quantity": 4}
        marker = {"row": 11, "marker": "fridge_side_panel_extras"}
        side_extra = {"row": 12, "item": "Шкант", "quantity": 2}
        customer_item = {
            "row": 17,
            "item": "Пуш",
            "quantity": 1,
            "external": True,
        }
        rows_by_sheet = {
            ("lower", "Холодильник"): [base, marker, side_extra, customer_item],
        }
        family = BUILD_SEED.FAMILIES["lower-fridge"]

        plain = BUILD_SEED.module_hardware_rows(
            family,
            {"variant": {"fridgeSidePanel": "none"}},
            rows_by_sheet,
        )
        with_side_panel = BUILD_SEED.module_hardware_rows(
            family,
            {"variant": {"fridgeSidePanel": "left"}},
            rows_by_sheet,
        )

        self.assertEqual([entry["row"] for entry in plain], [2, 17])
        self.assertEqual([entry["row"] for entry in with_side_panel], [2, 12, 17])


class FormulaCacheTests(unittest.TestCase):
    def test_missing_formula_cache_is_rejected_with_cell_address(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            workbook_path = Path(temporary_directory) / "formula-without-cache.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Корона"
            sheet["J2"] = "=2.7*0.1"
            workbook.save(workbook_path)

            cached_workbook = load_workbook(workbook_path, read_only=True, data_only=True)
            formula_workbook = load_workbook(workbook_path, read_only=True, data_only=False)
            try:
                with self.assertRaisesRegex(ValueError, r"Корона!J2"):
                    BUILD_SEED.cached_formula_value(
                        cached_workbook["Корона"]["J2"],
                        formula_workbook["Корона"]["J2"],
                        "Корона",
                    )
            finally:
                cached_workbook.close()
                formula_workbook.close()

    def test_missing_hardware_formula_cache_stops_import(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            workbook_path = Path(temporary_directory) / "Нижний ряд.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Угловые мойки"
            sheet.append(["Фурнитура", "кол-во", "ед.", "Себестоимость", "Комментарий"])
            sheet.append(["Евровинт", "=32", "шт", 32, None])
            workbook.save(workbook_path)

            with self.assertRaisesRegex(
                ValueError,
                r"Нижний ряд\.xlsx:Угловые мойки!B2",
            ):
                BUILD_SEED.parse_hardware_sources([workbook_path])


if __name__ == "__main__":
    unittest.main()
