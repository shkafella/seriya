from __future__ import annotations

import argparse
import collections
import hashlib
import json
import math
import re
from dataclasses import dataclass
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


SOURCE_TABS = ("низ гола", "среднии", "Антресоль", "Корона", "Фальш панели")
SPREADSHEET_ID = "19mCVoEg77-RVP0nzLKURPjWgAWv_yLM3WIhi8AbRx4I"
SOURCE_RECORD_PREFIX = "catalog-source-"
RATES = {
    "body_rub_m2": Decimal("2050"),
    "facade_rub_m2": Decimal("2225"),
    "hdf_rub_m2": Decimal("600"),
    "agt_1_rub_m2": Decimal("3000"),
    "packaging_rub_m2": Decimal("350"),
    "customer_price_multiplier": Decimal("1.10"),
    "included_vat_fraction": Decimal(22) / Decimal(122),
    "cut_to_width_service_rub": Decimal("3000"),
}


@dataclass(frozen=True)
class Family:
    family_id: str
    public_name: str
    row: str
    hardware_source: str | None
    hardware_sheet: str | None
    cut_policy: str = "allowed"


FAMILIES: dict[str, Family] = {
    "lower-sink": Family("lower-sink", "Нижний модуль под мойку", "lower", "lower", "Мойка"),
    "lower-dishwasher": Family("lower-dishwasher", "Модуль для посудомоечной машины", "lower", "lower", "ПМ", "forbidden"),
    "lower-shelf": Family("lower-shelf", "Нижний модуль с полкой", "lower", "lower", "Полка"),
    "lower-removable-shelf": Family("lower-removable-shelf", "Нижний модуль со съёмной полкой", "lower", "lower", "Полка съемная"),
    "lower-corner-sink": Family("lower-corner-sink", "Угловой нижний модуль под мойку", "lower", "lower", "Угловые мойки"),
    "lower-corner-removable-shelf": Family("lower-corner-removable-shelf", "Угловой нижний модуль со съёмной полкой", "lower", "lower", "Угловой съемная полка"),
    "lower-fridge": Family("lower-fridge", "Пенал для холодильника", "lower", "lower", "Холодильник", "forbidden"),
    "lower-two-drawers": Family("lower-two-drawers", "Нижний модуль с двумя ящиками", "lower", "lower", "2 ящика"),
    "lower-three-drawers": Family("lower-three-drawers", "Нижний модуль с тремя ящиками", "lower", "lower", "3 ящика"),
    "lower-hidden-drawers": Family("lower-hidden-drawers", "Нижний модуль со скрытой парой ящиков", "lower", "lower", "Скрытая пара"),
    "lower-oven": Family("lower-oven", "Нижний модуль для духовки", "lower", "lower", "Духовка", "forbidden"),
    "lower-bottle": Family("lower-bottle", "Бутылочница", "lower", "lower", "Бутылочница", "forbidden"),
    "wall-aventos": Family("wall-aventos", "Верхний модуль с подъёмником Aventos HF", "wall", "wall", "Полка Авентос HF"),
    "wall-klok": Family("wall-klok", "Верхний модуль с подъёмником KLOK", "wall", "wall", "Полка Клок или Газлифт"),
    "wall-shelf": Family("wall-shelf", "Верхний модуль с распашным фасадом", "wall", "wall", "Полка Распашной фасад"),
    "wall-removable-shelf": Family("wall-removable-shelf", "Верхний модуль со съёмной полкой", "wall", "wall", "Съемная полка Распашной фасад"),
    "wall-drying": Family("wall-drying", "Верхний модуль для сушки", "wall", "wall", "Сушки", "forbidden"),
    "wall-drying-aventos": Family("wall-drying-aventos", "Верхний модуль для сушки с Aventos HF", "wall", "wall", "Сушки Авентос HF", "forbidden"),
    "wall-hood": Family("wall-hood", "Верхний модуль для вытяжки", "wall", "wall", "Вытяжка", "forbidden"),
    "wall-corner": Family("wall-corner", "Угловой верхний модуль", "wall", "wall", "Угловые"),
    "mezzanine": Family("mezzanine", "Антресоль", "mezzanine", None, None),
    "cornice": Family("cornice", "Корона", "accessory", None, None, "forbidden"),
    "deep-cornice": Family("deep-cornice", "Глубокая корона", "accessory", None, None, "forbidden"),
    "plinth": Family("plinth", "Цоколь", "accessory", None, None, "forbidden"),
    "false-panel": Family("false-panel", "Фальш-панель", "accessory", None, None, "forbidden"),
}

TWO_FRONT_FAMILIES = {
    "lower-sink",
    "lower-shelf",
    "lower-removable-shelf",
    "wall-shelf",
    "wall-removable-shelf",
    "wall-drying",
}

def decimal(value: Any) -> Decimal:
    if value in (None, "", "нет"):
        return Decimal(0)
    return Decimal(str(value))


def decimal_number(value: Decimal) -> int | float:
    value = value.normalize()
    if value == value.to_integral():
        return int(value)
    return float(value)


def cached_formula_value(cached_cell: Any, formula_cell: Any, sheet_name: str) -> Any:
    formula_value = formula_cell.value
    is_formula = formula_cell.data_type == "f" or (
        isinstance(formula_value, str) and formula_value.startswith("=")
    )
    if is_formula and cached_cell.value is None:
        raise ValueError(
            "Нет сохранённого результата Excel-формулы в "
            f"{sheet_name}!{formula_cell.coordinate}. "
            "Откройте исходную таблицу в Excel или LibreOffice, пересчитайте формулы и сохраните файл."
        )
    return cached_cell.value


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def normalize_space(value: str) -> str:
    return re.sub(r"\s+", " ", value.strip())


def load_source_map(path: Path) -> tuple[dict[str, Any], bool]:
    if not path.exists():
        return {
            "schemaVersion": 1,
            "spreadsheetId": SPREADSHEET_ID,
            "records": [],
        }, True

    source_map = json.loads(path.read_text(encoding="utf-8"))
    if source_map.get("schemaVersion") != 1:
        raise ValueError("Неподдерживаемая версия source-map")
    if source_map.get("spreadsheetId") != SPREADSHEET_ID:
        raise ValueError("Source-map относится к другой Google-таблице")
    records = source_map.get("records")
    if not isinstance(records, list):
        raise ValueError("Source-map не содержит список records")
    record_ids = [record.get("sourceRecordId") for record in records]
    if any(not isinstance(record_id, str) or not record_id for record_id in record_ids):
        raise ValueError("Каждая запись source-map должна иметь sourceRecordId")
    if len(record_ids) != len(set(record_ids)):
        raise ValueError("Source-map содержит повторяющиеся sourceRecordId")
    return source_map, False


def next_source_record_id(source_map: dict[str, Any]) -> str:
    used_numbers = []
    for record in source_map["records"]:
        record_id = record["sourceRecordId"]
        match = re.fullmatch(rf"{re.escape(SOURCE_RECORD_PREFIX)}(\d+)", record_id)
        if match:
            used_numbers.append(int(match.group(1)))
    return f"{SOURCE_RECORD_PREFIX}{max(used_numbers, default=0) + 1:04d}"


def source_locator(row: dict[str, Any]) -> dict[str, Any]:
    return {
        "sheet": row["sourceSheet"],
        "row": row["sourceRow"],
        "sourceName": row["sourceName"],
        "sourceSku": row["sourceSku"],
    }


def assign_sheet_source_record_ids(
    source_rows: list[dict[str, Any]],
    source_map: dict[str, Any],
    allow_new_records: bool,
) -> set[str]:
    sheet_records = [record for record in source_map["records"] if record.get("kind") == "sheet-row"]
    assigned_ids: set[str] = set()

    for row in source_rows:
        locator = source_locator(row)
        available = [record for record in sheet_records if record["sourceRecordId"] not in assigned_ids]
        exact = [record for record in available if record.get("locator") == locator]
        if exact:
            candidates = exact
        else:
            candidates = []
            for record in available:
                previous = record.get("locator", {})
                if previous.get("sheet") != locator["sheet"]:
                    continue
                same_mutable_fields = sum(
                    previous.get(field) == locator[field]
                    for field in ("row", "sourceName", "sourceSku")
                )
                if same_mutable_fields >= 2:
                    candidates.append(record)

        if len(candidates) > 1:
            candidate_ids = ", ".join(record["sourceRecordId"] for record in candidates)
            raise ValueError(
                f"Неоднозначная source-map привязка для {locator}: {candidate_ids}. "
                "Укажите существующий sourceRecordId вручную в source-map."
            )
        if candidates:
            record = candidates[0]
            record["locator"] = locator
        elif allow_new_records:
            record = {
                "sourceRecordId": next_source_record_id(source_map),
                "kind": "sheet-row",
                "locator": locator,
            }
            source_map["records"].append(record)
            sheet_records.append(record)
        else:
            raise ValueError(
                f"Новая или одновременно переименованная и перемещённая строка {locator}. "
                "Добавьте явную привязку к sourceRecordId либо подтвердите добавление новой записи."
            )

        row["sourceRecordId"] = record["sourceRecordId"]
        assigned_ids.add(record["sourceRecordId"])

    stale_ids = {
        record["sourceRecordId"]
        for record in sheet_records
        if record["sourceRecordId"] not in assigned_ids
    }
    if stale_ids:
        raise ValueError(
            "Source-map содержит исчезнувшие строки: " + ", ".join(sorted(stale_ids))
        )
    return assigned_ids


def assign_derived_source_record_id(
    source_map: dict[str, Any],
    derived_key: str,
    allow_new_records: bool,
) -> str:
    candidates = [
        record
        for record in source_map["records"]
        if record.get("kind") == "derived" and record.get("derivedKey") == derived_key
    ]
    if len(candidates) > 1:
        raise ValueError(f"Source-map дублирует derivedKey {derived_key}")
    if candidates:
        return candidates[0]["sourceRecordId"]
    if not allow_new_records:
        raise ValueError(
            f"Source-map не содержит производную запись {derived_key}; добавьте её явно"
        )
    record = {
        "sourceRecordId": next_source_record_id(source_map),
        "kind": "derived",
        "derivedKey": derived_key,
    }
    source_map["records"].append(record)
    return record["sourceRecordId"]


def canonical_item_name(value: str) -> str:
    normalized = normalize_space(value).replace("ё", "е").lower()
    aliases = {
        "заглушки для евро винта": "Заглушка евровинта",
        "заглушка евровинта": "Заглушка евровинта",
        "рафикс": "Рафикс",
        "комлекты ящиков": "Комплекты ящиков",
    }
    if normalized in aliases:
        return aliases[normalized]
    return normalize_space(value)


def is_numeric(value: Any) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(float(value))


def hardware_source_key(path: Path) -> str:
    name = path.name.lower()
    if "антресоль" in name:
        return "mezzanine"
    if "корона" in name:
        return "crown"
    if "нижний" in name:
        return "lower"
    if "средний" in name:
        return "wall"
    raise ValueError(f"Неизвестный файл крепежа: {path}")


def family_for(sheet: str, sku: str, source_name: str) -> Family:
    compact = sku.replace(" ", "")
    if sheet == "низ гола":
        if compact.startswith("НГСП"):
            return FAMILIES["lower-removable-shelf"]
        if compact.startswith("НГПм"):
            return FAMILIES["lower-dishwasher"]
        if compact.startswith("НГП"):
            return FAMILIES["lower-shelf"]
        if compact.startswith("НГУлМ") or compact.startswith("НГУпМ"):
            return FAMILIES["lower-corner-sink"]
        if compact.startswith("НГУлСП") or compact.startswith("НГУпСП"):
            return FAMILIES["lower-corner-removable-shelf"]
        if compact.startswith("НГМ"):
            return FAMILIES["lower-sink"]
        if compact.startswith("НГХ"):
            return FAMILIES["lower-fridge"]
        if compact.startswith("НГСК") and compact.endswith("2"):
            return FAMILIES["lower-two-drawers"]
        if compact.startswith("НГСК") and compact.endswith("3"):
            return FAMILIES["lower-three-drawers"]
        if compact.startswith("НГСп") or compact.startswith("НГСП"):
            return FAMILIES["lower-hidden-drawers"]
        if compact == "НГДСК":
            return FAMILIES["lower-oven"]
        if compact.startswith("НГБ"):
            return FAMILIES["lower-bottle"]
    if sheet == "среднии":
        if compact.startswith("СДПА"):
            return FAMILIES["wall-aventos"]
        if compact.startswith("СДК"):
            return FAMILIES["wall-klok"]
        if compact.startswith("СДСП"):
            return FAMILIES["wall-removable-shelf"]
        if compact.startswith("СДСУА"):
            return FAMILIES["wall-drying-aventos"]
        if compact.startswith("СДСУ"):
            return FAMILIES["wall-drying"]
        if compact.startswith("СДУ"):
            return FAMILIES["wall-corner"]
        if "выт" in source_name.lower():
            return FAMILIES["wall-hood"]
        if compact.startswith("СДП"):
            return FAMILIES["wall-shelf"]
    if sheet == "Антресоль":
        return FAMILIES["mezzanine"]
    if sheet == "Корона":
        if sku.startswith("Корона Г"):
            return FAMILIES["deep-cornice"]
        if sku.startswith("Корона"):
            return FAMILIES["cornice"]
        if sku.startswith("Цоколь"):
            return FAMILIES["plinth"]
    if sheet == "Фальш панели":
        return FAMILIES["false-panel"]
    raise ValueError(f"Не определена семья: {sheet=} {sku=} {source_name=}")


def parse_hardware_sources(paths: Iterable[Path]) -> tuple[dict[tuple[str, str], list[dict[str, Any]]], dict[str, Decimal]]:
    rows_by_sheet: dict[tuple[str, str], list[dict[str, Any]]] = {}
    included_unit_rates: dict[str, list[Decimal]] = collections.defaultdict(list)

    for path in paths:
        source_key = hardware_source_key(path)
        workbook = load_workbook(path, read_only=True, data_only=True)
        formulas_workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            for sheet in workbook.worksheets:
                formulas_sheet = formulas_workbook[sheet.title]
                source_label = f"{path.name}:{sheet.title}"

                def cell_value(row_number: int, column: int) -> Any:
                    return cached_formula_value(
                        sheet.cell(row_number, column),
                        formulas_sheet.cell(row_number, column),
                        source_label,
                    )

                headers = {}
                for column, _cell in enumerate(
                    next(sheet.iter_rows(min_row=1, max_row=1)),
                    start=1,
                ):
                    value = cell_value(1, column)
                    if isinstance(value, str):
                        headers[value] = column
                item_col = headers.get("Фурнитура")
                qty_col = headers.get("кол-во")
                cost_col = headers.get("Себестоимость")
                parsed: list[dict[str, Any]] = []
                if item_col and qty_col and cost_col:
                    for row_number in range(2, (sheet.max_row or 0) + 1):
                        raw_item = cell_value(row_number, item_col)
                        if not isinstance(raw_item, str) or not raw_item.strip():
                            continue
                        if raw_item.strip().lower().startswith("на нгхф"):
                            parsed.append({"row": row_number, "marker": "fridge_side_panel_extras"})
                            continue
                        qty = cell_value(row_number, qty_col)
                        unit = cell_value(row_number, qty_col + 1)
                        raw_cost = cell_value(row_number, cost_col)
                        trailing = [
                            cell_value(row_number, column)
                            for column in range(cost_col + 1, min(cost_col + 3, (sheet.max_column or cost_col)) + 1)
                        ]
                        note_text = " ".join(str(value) for value in trailing if value not in (None, ""))
                        external = (
                            "покупают сами" in note_text.lower()
                            or (isinstance(raw_cost, str) and "покупают сами" in raw_cost.lower())
                            or "авентос" in raw_item.lower()
                            or (not is_numeric(raw_cost) and not is_numeric(qty))
                        )
                        item = canonical_item_name(raw_item)
                        entry = {
                            "row": row_number,
                            "item": item,
                            "source_item": normalize_space(raw_item),
                            "quantity": qty,
                            "unit": unit,
                            "source_cost": raw_cost,
                            "external": external,
                            "note": note_text or None,
                        }
                        if is_numeric(qty) and is_numeric(raw_cost) and not external and float(qty) != 0:
                            unit_rate = decimal(raw_cost) / decimal(qty)
                            entry["source_unit_rate"] = unit_rate
                            included_unit_rates[item].append(unit_rate)
                        parsed.append(entry)
                rows_by_sheet[(source_key, sheet.title)] = parsed
        finally:
            workbook.close()
            formulas_workbook.close()

    max_rates = {item: max(rates) for item, rates in included_unit_rates.items()}
    return rows_by_sheet, max_rates


def module_hardware_rows(
    family: Family,
    module: dict[str, Any],
    rows_by_sheet: dict[tuple[str, str], list[dict[str, Any]]],
) -> list[dict[str, Any]]:
    if family.family_id == "false-panel":
        return []
    if family.family_id == "mezzanine":
        sheet_name = "Антресоль КЛОК+ПУШ Г575" if module["dimensionsMm"]["depth"] == 575 else "Антресоль КЛОК+ПУШ Г320"
        return list(rows_by_sheet[("mezzanine", sheet_name)])
    if family.family_id in {"cornice", "deep-cornice", "plinth"}:
        rows = rows_by_sheet[("crown", "Антресоль КЛОК+ПУШ Г320")]
        row_ranges = {
            "cornice": range(2, 3),
            "deep-cornice": range(9, 11),
            "plinth": range(18, 19),
        }
        return [entry for entry in rows if entry.get("row") in row_ranges[family.family_id]]
    if not family.hardware_source or not family.hardware_sheet:
        return []
    rows = list(rows_by_sheet[(family.hardware_source, family.hardware_sheet)])
    if family.family_id == "lower-fridge":
        side_panel = module["variant"].get("fridgeSidePanel")
        return [
            entry
            for entry in rows
            if not entry.get("marker")
            and (
                entry.get("row", 0) < 11
                or entry.get("external")
                or (
                    side_panel in {"left", "right"}
                    and 11 < entry.get("row", 0) < 15
                )
            )
        ]
    return [entry for entry in rows if not entry.get("marker")]


def add_quantity(target: dict[str, dict[str, Any]], item: str, quantity: Decimal, unit: str = "шт") -> None:
    if item not in target:
        target[item] = {"item": item, "quantity": Decimal(0), "unit": unit}
    target[item]["quantity"] += quantity


def calculate_hardware(
    family: Family,
    module: dict[str, Any],
    rows_by_sheet: dict[tuple[str, str], list[dict[str, Any]]],
    max_rates: dict[str, Decimal],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], Decimal]:
    included: dict[str, dict[str, Any]] = {}
    supplied: dict[str, dict[str, Any]] = {}
    two_front = module["variant"].get("frontCount") == 2

    for entry in module_hardware_rows(family, module, rows_by_sheet):
        item = entry.get("item")
        qty = entry.get("quantity")
        if not item:
            continue
        if entry.get("external"):
            if is_numeric(qty):
                add_quantity(supplied, item, decimal(qty), entry.get("unit") or "шт")
                if (
                    "петл" in item.lower()
                    and family.family_id in TWO_FRONT_FAMILIES
                    and two_front
                ):
                    add_quantity(supplied, item, decimal(qty), entry.get("unit") or "шт")
            else:
                supplied[item] = {
                    "item": item,
                    "quantity": None,
                    "unit": entry.get("unit"),
                    "sourceQuantity": qty,
                }
            continue
        if is_numeric(qty) and item in max_rates:
            add_quantity(included, item, decimal(qty), entry.get("unit") or "шт")

    width = module["dimensionsMm"]["width"]
    if family.family_id in {"lower-sink", "lower-shelf", "lower-removable-shelf"} and two_front:
        add_quantity(included, "Саморез 4х16мм", Decimal(8))
    if family.family_id in {"lower-sink", "lower-shelf", "lower-removable-shelf"} and width >= 800:
        add_quantity(included, "Опора h=95-115мм", Decimal(2))
        add_quantity(included, "Саморез 4х16мм", Decimal(8))
    if family.family_id in {"wall-shelf", "wall-removable-shelf", "wall-drying"} and two_front:
        add_quantity(included, "Саморез 4х16мм", Decimal(8))
    if family.family_id == "mezzanine" and width >= 750:
        add_quantity(included, "Саморез 4х16мм", Decimal(9))
        add_quantity(supplied, "Петля накладная + ответка крест", Decimal(1))
        add_quantity(supplied, "KLOK", Decimal(1))
    if family.family_id == "wall-klok" and width >= 750:
        add_quantity(included, "Саморез 4х16мм", Decimal(18))
        add_quantity(supplied, "Петля накладная + ответка крест", Decimal(2))
        add_quantity(supplied, "KLOK", Decimal(2))

    breakdown: list[dict[str, Any]] = []
    total = Decimal(0)
    for item in sorted(included):
        entry = included[item]
        unit_rate = max_rates[item]
        line_total = entry["quantity"] * unit_rate
        total += line_total
        breakdown.append(
            {
                "item": item,
                "quantity": decimal_number(entry["quantity"]),
                "unit": entry["unit"],
                "unitRateRub": decimal_number(unit_rate),
                "lineTotalRub": decimal_number(line_total),
            }
        )

    supplied_rows = []
    for item in sorted(supplied):
        entry = supplied[item]
        supplied_rows.append(
            {
                "item": item,
                "quantity": None if entry["quantity"] is None else decimal_number(entry["quantity"]),
                "unit": entry.get("unit"),
                "sourceQuantity": entry.get("sourceQuantity"),
                "note": "Покупает клиент; не входит в цену и поставку.",
            }
        )
    return breakdown, supplied_rows, total


def front_count(family_id: str, sku: str, width: int) -> int | None:
    if family_id not in TWO_FRONT_FAMILIES:
        return None
    if sku.replace(" ", "").endswith("601"):
        return 2
    if family_id in {"lower-sink", "lower-shelf", "lower-removable-shelf", "wall-shelf", "wall-removable-shelf"} and width > 600:
        return 2
    if family_id == "wall-drying" and width > 600:
        return 2
    return 1


def corrected_identity(row: dict[str, Any]) -> tuple[str, int, list[str]]:
    family_id = row["family"].family_id
    sku = row["sourceSku"]
    width = int(row["width"])
    corrections: list[str] = []
    compact = sku.replace(" ", "")

    if family_id == "lower-removable-shelf":
        original = int(re.search(r"(\d+)$", compact).group(1))
        if original >= 650:
            corrected = 601 if original == 650 else original - 50
            sku = f"НГСП {corrected}"
            width = 600 if corrected == 601 else corrected
            corrections.append("D-58: исправлена сдвинутая последовательность НГСП")
    elif family_id in {"lower-sink", "lower-shelf", "wall-shelf", "wall-removable-shelf"}:
        match = re.search(r"(\d+)$", compact)
        if match:
            code = int(match.group(1))
            expected_width = 600 if code == 601 else code
            source_name_match = re.search(r"(\d+)\s*$", row["sourceName"])
            source_name_width = int(source_name_match.group(1)) if source_name_match else None
            if width != expected_width or source_name_width != expected_width:
                width = expected_width
                if family_id == "lower-sink":
                    corrections.append("D-37: ширина приведена к подтверждённой конструкции 600/601 и последовательности SKU")
                else:
                    corrections.append("Подтверждено владельцем: исправлена сдвинутая последовательность")
    if family_id == "wall-hood":
        sku = "СДВ 600"
        corrections.append("D-54: вытяжке назначен уникальный SKU СДВ 600")
    return sku, width, corrections


def variant_for(family: Family, sku: str, width: int, depth: int) -> dict[str, Any]:
    compact = sku.replace(" ", "")
    variant: dict[str, Any] = {}
    count = front_count(family.family_id, sku, width)
    if count:
        variant["frontCount"] = count
    if family.family_id in {"lower-corner-sink", "lower-corner-removable-shelf", "wall-corner"}:
        variant["openingSide"] = "left" if "л" in compact.lower() else "right"
    if family.family_id == "lower-fridge":
        variant["fridgeSidePanel"] = "left" if compact.endswith("Фл") else "right" if compact.endswith("Фп") else "none"
    if family.family_id == "mezzanine":
        variant["depth"] = depth
    return variant


def public_name(family: Family, module: dict[str, Any]) -> str:
    dimensions = module["dimensionsMm"]
    variant = module["variant"]
    if family.family_id == "false-panel":
        return f"{family.public_name} {dimensions['height']}×{dimensions['width']} мм"
    if family.family_id in {"cornice", "deep-cornice", "plinth"}:
        return f"{family.public_name} {dimensions['width']} мм"
    parts = [family.public_name, f"{dimensions['width']} мм"]
    if variant.get("frontCount") in {1, 2} and dimensions["width"] == 600:
        parts.append(f"{variant['frontCount']} фасад" if variant["frontCount"] == 1 else "2 фасада")
    if variant.get("openingSide"):
        parts.append("левый" if variant["openingSide"] == "left" else "правый")
    if variant.get("fridgeSidePanel") in {"left", "right"}:
        parts.append("фальш-панель слева" if variant["fridgeSidePanel"] == "left" else "фальш-панель справа")
    if family.family_id == "mezzanine":
        parts.append(f"глубина {dimensions['depth']} мм")
    return ", ".join(parts)


def round_ten(value: Decimal) -> Decimal:
    return (value / Decimal(10)).quantize(Decimal("1"), rounding=ROUND_HALF_UP) * Decimal(10)


def module_calculated_ex_vat(module: dict[str, Any]) -> Decimal:
    return (
        decimal(module["consumptionM2"]["body"]) * RATES["body_rub_m2"]
        + decimal(module["consumptionM2"]["facade"]) * RATES["facade_rub_m2"]
        + decimal(module["consumptionM2"]["hdf"]) * RATES["hdf_rub_m2"]
        + decimal(module["fasteners"]["chargeRub"])
    )


def build(args: argparse.Namespace) -> None:
    catalog_path = Path(args.catalog)
    hardware_paths = [Path(raw) for raw in args.hardware]
    output_path = Path(args.output)
    report_path = Path(args.report)
    example_path = Path(args.example)
    source_map_path = Path(args.source_map)
    source_map, source_map_was_created = load_source_map(source_map_path)
    allow_new_source_records = source_map_was_created or args.allow_new_source_records

    catalog_values = load_workbook(catalog_path, read_only=True, data_only=True)
    catalog_formulas = load_workbook(catalog_path, read_only=True, data_only=False)
    try:
        photo_sheet = catalog_values["Фото"]
        photo_by_source_name: dict[str, str] = {}
        for row in photo_sheet.iter_rows(min_row=2, max_col=2, values_only=True):
            name, url = row
            if isinstance(name, str) and isinstance(url, str) and url.startswith("http"):
                photo_by_source_name[normalize_space(name)] = url

        source_rows: list[dict[str, Any]] = []
        for sheet_name in SOURCE_TABS:
            values_sheet = catalog_values[sheet_name]
            formulas_sheet = catalog_formulas[sheet_name]
            values_rows = values_sheet.iter_rows(min_col=5, max_col=12)
            formulas_rows = formulas_sheet.iter_rows(min_col=5, max_col=12)
            for values_cells, formula_cells in zip(values_rows, formulas_rows, strict=True):
                source_name = values_cells[0].value
                source_sku = values_cells[1].value
                if not (isinstance(source_name, str) and isinstance(source_sku, str)):
                    continue
                if source_name.strip().lower() == "название" or source_sku.strip().lower() == "имя":
                    continue
                values = tuple(
                    cached_formula_value(values_cell, formula_cell, sheet_name)
                    for values_cell, formula_cell in zip(values_cells, formula_cells, strict=True)
                )
                source_name, source_sku, height, width, depth, body_m2, facade_m2, hdf_m2 = values
                family = family_for(sheet_name, normalize_space(source_sku), normalize_space(source_name))
                source_rows.append(
                    {
                        "sourceSheet": sheet_name,
                        "sourceRow": values_cells[0].row,
                        "sourceName": normalize_space(source_name),
                        "sourceSku": normalize_space(source_sku),
                        "height": height,
                        "width": width,
                        "depth": depth,
                        "bodyM2": body_m2,
                        "facadeM2": facade_m2,
                        "hdfM2": hdf_m2,
                        "family": family,
                    }
                )
    finally:
        catalog_values.close()
        catalog_formulas.close()

    assigned_source_record_ids = assign_sheet_source_record_ids(
        source_rows,
        source_map,
        allow_new_source_records,
    )

    rows_by_sheet, max_rates = parse_hardware_sources(hardware_paths)
    modules: list[dict[str, Any]] = []
    corrections: list[dict[str, Any]] = []

    for row in source_rows:
        family: Family = row["family"]
        sku, width, row_corrections = corrected_identity(row)
        height_source = row["height"]
        height = 2568 if isinstance(height_source, str) else int(height_source)
        depth = int(row["depth"])
        variant = variant_for(family, sku, width, depth)
        module = {
            "sourceRecordId": row["sourceRecordId"],
            "sku": sku,
            "familyId": family.family_id,
            "familyName": family.public_name,
            "row": family.row,
            "publicName": "",
            "dimensionsMm": {"height": height, "width": width, "depth": depth},
            "heightRule": "complete_kitchen_height_excluding_crown" if family.family_id == "lower-fridge" else "fixed",
            "consumptionM2": {
                "body": decimal_number(decimal(row["bodyM2"])),
                "facade": decimal_number(decimal(row["facadeM2"])),
                "hdf": decimal_number(decimal(row["hdfM2"])),
            },
            "variant": variant,
            "cutToWidth": {
                "policy": family.cut_policy,
                "maxReductionMm": 50 if family.cut_policy != "forbidden" else None,
                "stepMm": 1 if family.cut_policy != "forbidden" else None,
            },
            "fasteners": {},
            "sourceImageUrl": photo_by_source_name.get(row["sourceName"]),
            "source": {
                "spreadsheetId": "19mCVoEg77-RVP0nzLKURPjWgAWv_yLM3WIhi8AbRx4I",
                "sheet": row["sourceSheet"],
                "row": row["sourceRow"],
                "originalName": row["sourceName"],
                "originalSku": row["sourceSku"],
                "originalHeight": height_source,
                "originalWidth": row["width"],
            },
            "review": {
                "publicationStatus": "approved",
                "pending": [],
            },
        }
        if family.family_id == "false-panel":
            module["layoutSpanMm"] = depth
        module["publicName"] = public_name(family, module)
        if row_corrections:
            module["source"]["corrections"] = row_corrections
            corrections.append(
                {
                    "sourceSheet": row["sourceSheet"],
                    "sourceRow": row["sourceRow"],
                    "originalSku": row["sourceSku"],
                    "normalizedSku": sku,
                    "notes": row_corrections,
                }
            )
        breakdown, supplied, charge = calculate_hardware(family, module, rows_by_sheet, max_rates)
        module["fasteners"] = {
            "chargeRub": decimal_number(charge),
            "included": breakdown,
            "customerSupplied": supplied,
        }
        modules.append(module)

    drying_500 = next(module for module in modules if module["sku"] == "СДСУА 500")
    regular_drying_450 = next(module for module in modules if module["sku"] == "СДСУ 450")
    derived_450 = json.loads(json.dumps(drying_500, ensure_ascii=False))
    derived_450["sourceRecordId"] = assign_derived_source_record_id(
        source_map,
        "wall-drying-aventos-450",
        allow_new_source_records,
    )
    assigned_source_record_ids.add(derived_450["sourceRecordId"])
    derived_450["sku"] = "СДСУА 450"
    derived_450["dimensionsMm"]["width"] = 450
    derived_450["consumptionM2"] = dict(regular_drying_450["consumptionM2"])
    derived_450["publicName"] = "Верхний модуль для сушки с Aventos HF, 450 мм"
    derived_450["sourceImageUrl"] = drying_500["sourceImageUrl"]
    derived_450["source"] = {
        "derivedFrom": [
            {"sku": "СДСУА 500", "purpose": "семья и конструкция Aventos HF"},
            {"sku": "СДСУ 450", "purpose": "утверждённая формула расхода для ширины 450 мм"},
        ],
        "corrections": ["D-39 и D-47: восстановлен отсутствующий вариант 450 мм"],
    }
    derived_450["review"] = {
        "publicationStatus": "approved",
        "pending": [],
    }
    breakdown, supplied, charge = calculate_hardware(FAMILIES["wall-drying-aventos"], derived_450, rows_by_sheet, max_rates)
    derived_450["fasteners"] = {
        "chargeRub": decimal_number(charge),
        "included": breakdown,
        "customerSupplied": supplied,
    }
    modules.append(derived_450)
    corrections.append(
        {
            "sourceSheet": "среднии",
            "sourceRow": 89,
            "originalSku": "СДСУА 500",
            "normalizedSku": "СДСУА 450 + СДСУА 500",
            "notes": ["D-39 и D-47: разделены варианты 450 и 500 мм"],
        }
    )

    modules.sort(key=lambda item: (item["row"], item["familyId"], item["dimensionsMm"]["width"], item["sku"]))
    unassigned_source_records = {
        record["sourceRecordId"]
        for record in source_map["records"]
        if record["sourceRecordId"] not in assigned_source_record_ids
    }
    if unassigned_source_records:
        raise ValueError(
            "Source-map содержит неиспользуемые записи: "
            + ", ".join(sorted(unassigned_source_records))
        )
    source_map["records"].sort(key=lambda record: record["sourceRecordId"])
    source_map_text = json.dumps(source_map, ensure_ascii=False, indent=2) + "\n"
    source_map_checksum_payload = json.dumps(
        source_map,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    )
    source_map_sha256 = hashlib.sha256(
        source_map_checksum_payload.encode("utf-8")
    ).hexdigest()
    counts = collections.Counter(module["row"] for module in modules)
    family_counts = collections.Counter(module["familyId"] for module in modules)
    unique_images = sorted({module["sourceImageUrl"] for module in modules if module["sourceImageUrl"]})
    missing_images = [module["sku"] for module in modules if not module["sourceImageUrl"]]

    conflicts = []
    all_observed: dict[str, set[Decimal]] = collections.defaultdict(set)
    for entries in rows_by_sheet.values():
        for entry in entries:
            if entry.get("source_unit_rate") is not None:
                all_observed[entry["item"]].add(entry["source_unit_rate"])
    for item, observed in sorted(all_observed.items()):
        if len(observed) > 1:
            conflicts.append(
                {
                    "item": item,
                    "observedUnitRatesRub": [decimal_number(value) for value in sorted(observed)],
                    "selectedUnitRateRub": decimal_number(max(observed)),
                }
            )

    source_files = [catalog_path, *hardware_paths]
    catalog = {
        "schemaVersion": 1,
        "status": "approved_stage_1",
        "generatedAt": "2026-10-21",
        "source": {
            "googleSpreadsheetId": SPREADSHEET_ID,
            "googleSpreadsheetUrl": f"https://docs.google.com/spreadsheets/d/{SPREADSHEET_ID}/edit",
            "googleModifiedAt": args.google_modified_at,
            "sourceMap": {
                "name": source_map_path.name,
                "sha256": source_map_sha256,
                "recordCount": len(source_map["records"]),
            },
            "files": [
                {"name": path.name, "sha256": sha256(path), "role": "catalog" if path == catalog_path else "fasteners"}
                for path in source_files
            ],
        },
        "commercialRules": {
            "currency": "RUB",
            "rates": {key: decimal_number(value) for key, value in RATES.items()},
            "bodyEdgingIncludedInBodyRate": True,
            "facadeEdgingIncludedInFacadeRate": True,
            "countertopIncluded": False,
            "assemblyIncluded": False,
            "customerSuppliedComponentsIncluded": False,
            "finalCustomerTotalRoundingRub": 10,
            "vatDisplayAccountantApproved": True,
        },
        "materials": {
            "body": [
                {"id": "ldsp-tefi", "name": "ЛДСП Tefi", "rateRubM2": 2050},
                {"id": "ldsp-white", "name": "ЛДСП белый", "rateRubM2": 2050},
                {"id": "lmdf-ultradecor-white", "name": "ЛМДФ Ultrаdecor белый", "rateRubM2": 2050},
                {"id": "lmdf-ultradecor-cashmere", "name": "ЛМДФ Ultrаdecor кашемир", "rateRubM2": 2050},
                {"id": "lmdf-ultradecor-graphite", "name": "ЛМДФ Ultrаdecor серый графит", "rateRubM2": 2050},
            ],
            "facade": [
                *[
                    {"id": f"ultradecor-{index}", "name": f"Ultrаdecor цвет {index}", "rateRubM2": 2225}
                    for index in range(1, 7)
                ],
                {"id": "agt-1", "name": "AGT 1", "rateRubM2": 3000},
            ],
            "hdf": [{"id": "hdf-white", "name": "ХДФ белый", "rateRubM2": 600}],
        },
        "fastenerUnitRates": [
            {"item": item, "rateRub": decimal_number(rate)} for item, rate in sorted(max_rates.items())
        ],
        "families": [
            {
                "id": family.family_id,
                "publicName": family.public_name,
                "row": family.row,
                "cutPolicy": family.cut_policy,
            }
            for family in FAMILIES.values()
        ],
        "modules": modules,
    }

    report = {
        "status": "valid_launch_rules_approved",
        "checks": {
            "moduleCount": len(modules),
            "expectedModuleCount": 225,
            "rowCounts": dict(sorted(counts.items())),
            "familyCounts": dict(sorted(family_counts.items())),
            "uniqueSkuCount": len({module["sku"] for module in modules}),
            "uniqueSourceRecordIdCount": len({module["sourceRecordId"] for module in modules}),
            "sourceMapRecordCount": len(source_map["records"]),
            "uniqueImageUrlCount": len(unique_images),
            "missingImageSkus": missing_images,
            "negativeFastenerCharges": [
                module["sku"] for module in modules if module["fasteners"]["chargeRub"] < 0
            ],
        },
        "corrections": corrections,
        "fastenerUnitRateConflicts": conflicts,
        "pendingOwnerReviews": [],
        "pendingLaunchReviews": [],
    }

    sku_map = {module["sku"]: module for module in modules}
    example_skus = [
        "НГБ 300",
        "НГМ 600",
        "НГДСК",
        "НГПм 600",
        "НГСК 602",
        "СДК 600",
        "СДП 600",
        "СДСУ 600",
        "СДПА 500",
        "СДК 400",
        "А 600",
        "А 600",
        "А 600",
        "А 500",
        "А 400",
    ]
    selected = [sku_map[sku] for sku in example_skus]
    material_total = sum(module_calculated_ex_vat(module) for module in selected)
    packaging_total = sum(
        decimal(module["consumptionM2"]["body"]) * RATES["packaging_rub_m2"]
        for module in selected
    )
    ex_vat = material_total + packaging_total
    customer_total = round_ten(ex_vat * RATES["customer_price_multiplier"])
    included_vat = (customer_total * RATES["included_vat_fraction"]).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    false_panel = sku_map["ФП 350*341"]
    plinth_product = sku_map["Цоколь 2700"]
    crown_product = sku_map["Корона 2700"]
    plinth_covered_run = 5400 + false_panel["layoutSpanMm"]
    crown_covered_run = 2700 + 600
    plinth_quantity = math.ceil(plinth_covered_run / plinth_product["dimensionsMm"]["width"])
    crown_quantity = math.ceil(crown_covered_run / crown_product["dimensionsMm"]["width"])
    false_panel_price = module_calculated_ex_vat(false_panel)
    plinth_unit_price = module_calculated_ex_vat(plinth_product)
    crown_unit_price = module_calculated_ex_vat(crown_product)
    example = {
        "schemaVersion": 1,
        "status": "approved_reference",
        "name": "Контрольная прямая кухня 2700 мм",
        "countertopThicknessMm": 38,
        "completeHeightMm": 2568,
        "wallLengthMm": 2700,
        "options": {"facades": True, "packaging": True, "cutToWidth": []},
        "materials": {"body": "ldsp-white", "facade": "ultradecor-1", "hdf": "hdf-white"},
        "rows": {
            "lower": example_skus[:5],
            "wall": example_skus[5:10],
            "mezzanine": example_skus[10:],
        },
        "expected": {
            "rowWidthMm": {"lower": 2700, "wall": 2700, "mezzanine": 2700},
            "calculatedExVatRub": decimal_number(ex_vat),
            "packagingRub": decimal_number(packaging_total),
            "customerTotalRub": decimal_number(customer_total),
            "includedVatRub": decimal_number(included_vat),
        },
        "accessoryRegression": {
            "falsePanel": {
                "sku": false_panel["sku"],
                "layoutSpanMm": false_panel["layoutSpanMm"],
                "continuesPlinthFrontLine": True,
                "continuesCrownFrontLine": False,
                "calculatedExVatRub": decimal_number(false_panel_price),
            },
            "plinth": {
                "sku": plinth_product["sku"],
                "contributingSpansMm": {
                    "floorStandingFrontIncludingTallUnit": 5400,
                    "tallUnitIncludedInFloorStandingFront": 600,
                    "falsePanelContinuingFrontLine": false_panel["layoutSpanMm"],
                    "sideReturns": 0,
                },
                "coveredRunMm": plinth_covered_run,
                "stockPieceLengthMm": plinth_product["dimensionsMm"]["width"],
                "quantity": plinth_quantity,
                "unitCalculatedExVatRub": decimal_number(plinth_unit_price),
                "subtotalCalculatedExVatRub": decimal_number(plinth_unit_price * plinth_quantity),
            },
            "crown": {
                "sku": crown_product["sku"],
                "contributingSpansMm": {
                    "mezzanineFront": 2700,
                    "tallUnitTop": 600,
                    "falsePanelContinuingFrontLine": 0,
                    "sideReturns": 0,
                },
                "coveredRunMm": crown_covered_run,
                "stockPieceLengthMm": crown_product["dimensionsMm"]["width"],
                "quantity": crown_quantity,
                "unitCalculatedExVatRub": decimal_number(crown_unit_price),
                "subtotalCalculatedExVatRub": decimal_number(crown_unit_price * crown_quantity),
            },
        },
        "lead": {
            "company": "Контрольная мебельная компания",
            "contactName": "Тестовый клиент",
            "phone": "+7 900 000-00-00",
            "preferredMessenger": "telegram",
            "messengerContact": "@test",
            "comment": "Тестовая заявка; не отправлять.",
        },
    }

    for path in (source_map_path, output_path, report_path, example_path):
        path.parent.mkdir(parents=True, exist_ok=True)
    source_map_path.write_bytes(source_map_text.encode("utf-8"))
    output_path.write_text(json.dumps(catalog, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    example_path.write_text(json.dumps(example, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--catalog", required=True)
    parser.add_argument("--hardware", action="append", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--report", required=True)
    parser.add_argument("--example", required=True)
    parser.add_argument("--source-map", required=True)
    parser.add_argument("--allow-new-source-records", action="store_true")
    parser.add_argument("--google-modified-at", default="2026-10-20T11:29:29.114Z")
    return parser.parse_args()


if __name__ == "__main__":
    build(parse_args())
